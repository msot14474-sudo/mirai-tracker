from __future__ import annotations

import os
from pathlib import Path

from PIL import Image, ImageChops
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "build" / "role-workflows"
OUTPUT = BUILD / "Mirai_Tracker_Role_Workflows_and_Access.docx"
LOGO = ROOT / "apps" / "web" / "public" / "mirai-logo.jpg"
DIAGRAM_SOURCES = [
    Path("/tmp/mirai-gateway-check/role-diagram-1.png"),
    Path("/tmp/mirai-gateway-check/role-diagram-2.png"),
    Path("/tmp/mirai-gateway-check/role-diagram-3.png"),
]

NAVY = "173B7A"
BLUE = "2563EB"
PALE_BLUE = "EAF2FF"
PALE_TEAL = "E6F8F5"
PALE_GOLD = "FFF4D6"
PALE_PURPLE = "F0EAFE"
PALE_GRAY = "F4F7FB"
MID_GRAY = "667085"
LINE = "DCE3EC"
BLACK = "111827"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def crop_diagram(source: Path, destination: Path) -> None:
    image = Image.open(source).convert("RGB")
    background = Image.new("RGB", image.size, "white")
    diff = ImageChops.difference(image, background).convert("L")
    diff = diff.point(lambda px: 0 if px < 12 else 255)
    bbox = diff.getbbox()
    if bbox is None:
        image.save(destination)
        return
    left, top, right, bottom = bbox
    pad_x = max(36, int((right - left) * 0.025))
    pad_y = max(36, int((bottom - top) * 0.08))
    crop_box = (
        max(0, left - pad_x),
        max(0, top - pad_y),
        min(image.width, right + pad_x),
        min(image.height, bottom + pad_y),
    )
    image.crop(crop_box).save(destination, quality=95)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_section_title(doc: Document, kicker: str, title: str, intro: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(kicker.upper())
    set_run_font(r, size=9, bold=True, color=BLUE)
    r.font.letter_spacing = Pt(0.8)
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_after = Pt(6)
    h.add_run(title)
    if intro:
        p = doc.add_paragraph(intro)
        p.style = doc.styles["Subtitle"]
        p.paragraph_format.space_after = Pt(12)


def add_caption(doc: Document, label: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(label)
    set_run_font(r, size=9, bold=True, color=MID_GRAY)


def add_bullets(cell, items: list[str]) -> None:
    for item in items:
        p = cell.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.17)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=9.5, color=BLACK)


def add_role_card(doc: Document, role: str, scope: str, actions: list[str], sees: list[str], fill: str) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.8)
    table.columns[1].width = Inches(4.8)
    header = table.rows[0].cells
    merged = header[0].merge(header[1])
    set_cell_shading(merged, fill)
    set_cell_border(merged)
    set_cell_margins(merged, 130, 170, 130, 170)
    p = merged.paragraphs[0]
    r = p.add_run(role)
    set_run_font(r, size=14, bold=True, color=BLACK)
    r = p.add_run(f"  ·  {scope}")
    set_run_font(r, size=10, bold=False, color=MID_GRAY)

    for index, (heading, items) in enumerate((("Can do", actions), ("Can see", sees))):
        cell = table.rows[1].cells[index]
        set_cell_border(cell)
        set_cell_margins(cell, 120, 170, 120, 170)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        r = p.add_run(heading)
        set_run_font(r, size=10.5, bold=True, color=NAVY)
        add_bullets(cell, items)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(34)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string(MID_GRAY)

    for name, size in (("Heading 1", 24), ("Heading 2", 15), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header.distance = Inches(0.22)
    table = header.add_table(rows=1, cols=2, width=Inches(9.7))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(8.9)
    table.columns[1].width = Inches(0.8)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = left.add_run("MIRAI TRACKER  ·  ROLE WORKFLOWS")
    set_run_font(run, size=8, bold=True, color=NAVY)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = right.add_run("v1.0")
    set_run_font(run, size=8, color=MID_GRAY)

    footer = section.footer
    footer.distance = Inches(0.22)
    add_page_number(footer.paragraphs[0])


def build_document() -> Path:
    BUILD.mkdir(parents=True, exist_ok=True)
    cropped = []
    for index, source in enumerate(DIAGRAM_SOURCES, 1):
        destination = BUILD / f"role-diagram-{index}.png"
        crop_diagram(source, destination)
        cropped.append(destination)

    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Mirai Tracker Role Workflows and Access"
    props.subject = "Academic delivery flow, role responsibilities, permissions and notifications"
    props.author = "Mirai School of Technology"
    props.keywords = "Mirai Tracker, workflow, roles, permissions, academic delivery"

    # Cover
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_p.paragraph_format.space_after = Pt(16)
    logo_p.add_run().add_picture(str(LOGO), width=Inches(0.72))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("ACADEMIC OPERATING GUIDE")
    set_run_font(r, size=9.5, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(10)
    title.add_run("Role workflows\nand access")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle.add_run("How semester decisions move through campus delivery, faculty planning and the student view.")

    summary = doc.add_table(rows=1, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary.autofit = False
    for index, (number, label, fill) in enumerate(
        (("01", "Global academic setup", PALE_BLUE), ("02", "Campus delivery", PALE_TEAL), ("03", "Teaching execution", PALE_GOLD), ("04", "Read-only student view", PALE_PURPLE))
    ):
        cell = summary.cell(0, index)
        set_cell_shading(cell, fill)
        set_cell_border(cell, WHITE)
        set_cell_margins(cell, 150, 150, 150, 150)
        p = cell.paragraphs[0]
        r = p.add_run(number)
        set_run_font(r, size=12, bold=True, color=NAVY)
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=10, bold=True, color=BLACK)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Operating principle")
    set_run_font(r, size=12, bold=True, color=BLACK)
    p = doc.add_paragraph(
        "Authority flows downward by role and campus scope. Progress, exceptions and reports flow upward. Every person sees only the campuses, batches and subjects assigned to them."
    )
    p.paragraph_format.space_after = Pt(4)
    for text in (
        "Administrator / Program Manager — all campuses",
        "Campus Manager — assigned campus",
        "Faculty and Visiting Faculty — assigned subjects and classes",
        "Student — enrolled campus and batch, read only",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(text)

    doc.add_page_break()

    # End-to-end flow
    add_section_title(
        doc,
        "01 · End-to-end flow",
        "From semester setup to the student view",
        "The platform connects administrative planning, campus scheduling and classroom delivery in one controlled chain.",
    )
    picture = doc.add_picture(str(cropped[0]), width=Inches(9.55))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_alt_text(picture, "End-to-end academic flow", "Administrator creates the semester, campus manager builds the timetable, faculty plans and logs teaching, and the student view refreshes.")
    add_caption(doc, "Figure 1 · End-to-end academic delivery flow")
    p = doc.add_paragraph()
    r = p.add_run("What changes downstream: ")
    set_run_font(r, size=10, bold=True, color=NAVY)
    r = p.add_run("A calendar or exam-window change recalculates affected classes, planned dates, available teaching hours, progress and exam readiness before it reaches the student view.")
    set_run_font(r, size=10, color=BLACK)

    doc.add_page_break()

    # Hierarchy
    add_section_title(
        doc,
        "02 · Access hierarchy",
        "Authority follows campus and subject scope",
        "Each role receives the minimum scope needed to complete its part of the academic workflow.",
    )
    picture = doc.add_picture(str(cropped[1]), width=Inches(8.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_alt_text(picture, "Role access hierarchy", "Administrator assigns campus managers; campus managers assign faculty, visiting faculty and students; reports flow upward.")
    add_caption(doc, "Figure 2 · Role hierarchy and reporting paths")
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (heading, detail, fill) in enumerate(
        (
            ("Global", "Administrator and Program Manager configure all campuses and review cross-campus intelligence.", PALE_BLUE),
            ("Campus", "Campus Managers operate only inside their assigned campus and manage delivery.", PALE_TEAL),
            ("Assigned work", "Faculty see assigned teaching; students see enrolled subjects and outcomes.", PALE_GOLD),
        )
    ):
        cell = table.cell(0, index)
        set_cell_shading(cell, fill)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        r = p.add_run(heading)
        set_run_font(r, size=11, bold=True, color=BLACK)
        cell = table.cell(1, index)
        set_cell_border(cell)
        set_cell_margins(cell, 110, 140, 110, 140)
        p = cell.paragraphs[0]
        r = p.add_run(detail)
        set_run_font(r, size=9.5, color=BLACK)

    # Leadership roles
    add_section_title(
        doc,
        "03 · Role guide",
        "Leadership and campus operations",
        "Administrative roles establish the academic structure and keep delivery aligned when conditions change.",
    )
    add_role_card(
        doc,
        "Administrator / Program Manager",
        "All campuses",
        [
            "Create semesters and maintain calendars, holidays, working days and exam windows.",
            "Create subjects; manage users, campus managers and campus assignments.",
            "Review notifications, reports, exports, audit history and AI summaries.",
        ],
        [
            "Cross-campus progress and teaching delivery.",
            "Faculty activity and exam readiness.",
            "Campus, subject, batch and track analytics across the organisation.",
        ],
        PALE_BLUE,
    )
    add_role_card(
        doc,
        "Campus Manager",
        "Assigned campus",
        [
            "Create or link subjects, batches and people; assign faculty to subjects.",
            "Upload and edit the full timetable, alter a particular class, and add extra classes.",
            "Maintain permitted calendar changes and manage visiting-faculty per-class rates.",
        ],
        [
            "Campus subjects, timetable, lecture delivery and faculty activity.",
            "Progress, exam readiness, notifications and campus reports.",
            "Visiting-faculty class counts, approved payments and honorarium totals.",
        ],
        PALE_TEAL,
    )

    doc.add_page_break()

    # Teaching roles
    add_section_title(
        doc,
        "04 · Role guide",
        "Teaching and student experience",
        "Teaching roles convert assigned hours into lesson plans and recorded delivery; students receive the resulting read-only view.",
    )
    add_role_card(
        doc,
        "Faculty",
        "Assigned subjects only",
        [
            "Upload a partial or full-semester lesson plan and manage planned dates.",
            "Record daily lecture delivery and respond to schedule-change notifications.",
            "Use optional WhatsApp logging when that integration is enabled.",
        ],
        [
            "Assigned hours, upcoming classes, covered and pending topics.",
            "Subject progress, buffer ratio, exam readiness, AI insights and own reports.",
            "Notifications only for assigned subjects and teaching actions.",
        ],
        PALE_GOLD,
    )
    add_role_card(
        doc,
        "Visiting Faculty",
        "Assigned subjects and classes",
        [
            "Follow the Faculty workflow after receiving subject access.",
            "Record classes actually delivered for approval and payment.",
        ],
        [
            "Assigned hours, timetable and subject plan.",
            "Personal delivered-class count; honorarium is calculated per approved class.",
        ],
        PALE_GOLD,
    )
    add_role_card(
        doc,
        "Student",
        "Enrolled campus and batch · read only",
        ["No academic editing or access-management actions."],
        [
            "Enrolled subjects and assigned teachers.",
            "Covered and upcoming topics, exam dates and syllabus scope.",
        ],
        PALE_PURPLE,
    )

    doc.add_page_break()

    # Change flow
    add_section_title(
        doc,
        "05 · Change management",
        "Calendar changes trigger targeted updates",
        "The system identifies impact first, asks the responsible faculty member when a teaching decision is required, and then refreshes dependent views.",
    )
    picture = doc.add_picture(str(cropped[2]), width=Inches(9.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_alt_text(picture, "Calendar change and notification flow", "Calendar change finds affected classes, notifies assigned faculty when needed, updates the timetable and planned dates, then refreshes student and manager views.")
    add_caption(doc, "Figure 3 · Impact, decision and notification flow")
    steps = doc.add_table(rows=2, cols=4)
    steps.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (number, title, text_value) in enumerate(
        (
            ("1", "Detect", "Find every future class affected by the calendar or exam-window change."),
            ("2", "Decide", "Ask the assigned faculty member when topic order or recovery timing needs a choice."),
            ("3", "Re-align", "Update timetable dates, available hours and linked lesson-plan timing."),
            ("4", "Notify", "Show the result to the right faculty, manager, administrator and students."),
        )
    ):
        cell = steps.cell(0, index)
        set_cell_shading(cell, PALE_BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        r = p.add_run(f"{number}  {title}")
        set_run_font(r, size=10.5, bold=True, color=NAVY)
        cell = steps.cell(1, index)
        set_cell_border(cell)
        set_cell_margins(cell, 100, 120, 100, 120)
        p = cell.paragraphs[0]
        r = p.add_run(text_value)
        set_run_font(r, size=9.2, color=BLACK)

    doc.add_page_break()

    # Permission matrix
    add_section_title(
        doc,
        "06 · Permission reference",
        "What each role may access",
        "The interface and API apply the same role and scope checks. A visible control is never the security boundary.",
    )
    headers = ["Capability", "Admin", "Campus manager", "Faculty", "Student"]
    rows = [
        ("All-campus analytics", "Full", "—", "—", "—"),
        ("Assigned-campus operations", "Full", "Full", "—", "Read only"),
        ("Semester and calendar maintenance", "Full", "Assigned campus", "—", "—"),
        ("Subject and faculty assignment", "Full", "Assigned campus", "—", "—"),
        ("Timetable management", "Full", "Assigned campus", "View assigned", "View enrolled"),
        ("Lesson plans and lecture logs", "Review", "Review campus", "Own subjects", "View outcomes"),
        ("Reports", "All campuses", "Assigned campus", "Own teaching", "—"),
        ("Notifications", "Global / escalated", "Campus", "Assigned subjects", "Relevant results"),
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.0, 1.5, 2.1, 1.5, 1.5]
    for idx, (cell, heading) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell, 100, 110, 100, 110)
        p = cell.paragraphs[0]
        r = p.add_run(heading)
        set_run_font(r, size=9.5, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row_data):
            set_cell_shading(cells[col_index], WHITE if row_index % 2 == 0 else PALE_GRAY)
            set_cell_border(cells[col_index])
            set_cell_margins(cells[col_index], 82, 110, 82, 110)
            p = cells[col_index].paragraphs[0]
            r = p.add_run(value)
            set_run_font(r, size=9.2, bold=(col_index == 0), color=BLACK)

    h = doc.add_paragraph(style="Heading 2")
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(5)
    h.add_run("Permission principles")
    for text in (
        "Campus managers cannot manage another campus.",
        "Faculty cannot view or edit another faculty member’s subject unless separately assigned.",
        "Students cannot edit calendars, timetables, plans, logs, progress or user access.",
        "Changes to teaching history retain the responsible user and time for review.",
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(text)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
